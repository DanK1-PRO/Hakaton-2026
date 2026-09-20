"""Normalize source rows; preserve conditional routing without inventing conditions."""

import hashlib
import io
import json
from pathlib import Path
from zipfile import ZipFile
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
archive_path = ROOT / "Данные заказчика/Датасет.zip"
with ZipFile(archive_path) as archive:
    name = next(n for n in archive.namelist() if n.endswith(".xlsx"))
    raw = archive.read(name)
ws = openpyxl.load_workbook(io.BytesIO(raw), data_only=True).active
# Expand merged headers only, never data cells.
headers = [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)] for r in (1, 2, 3)]
for region in ws.merged_cells.ranges:
    if region.min_row <= 3:
        value = ws.cell(region.min_row, region.min_col).value
        for r in range(region.min_row, min(3, region.max_row) + 1):
            for c in range(region.min_col, region.max_col + 1):
                headers[r - 1][c - 1] = value
items = []
for row in range(4, ws.max_row + 1):
    values = [ws.cell(row, c).value for c in range(1, ws.max_column + 1)]
    if not values[10] or values[4] is None:
        continue
    routes = []
    for c in range(13, len(values)):
        if values[c] is None or not str(values[c]).strip():
            continue
        routes.append(
            {
                "column": c + 1,
                "service": str(headers[0][c] or ""),
                "variant": str(headers[1][c] or ""),
                "condition": str(headers[2][c] or ""),
                "value": str(values[c]).strip(),
            }
        )
    items.append(
        {
            "id": row,
            "external_code": str(values[4]),
            "name": str(values[10]).strip(),
            "features": [str(v) for v in values[6:10] if v is not None],
            "routes": routes,
            "source": {"file": name, "sheet": ws.title, "row": row, "sha256": hashlib.sha256(raw).hexdigest()},
        }
    )
out = ROOT / "data_derived/classifier"
out.mkdir(parents=True, exist_ok=True)
(out / "incident_types.json").write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Normalized {len(items)} classifier rows; conditional routes preserved")
by_code = {x["external_code"]: x for x in items}
cases = [
    (
        "water",
        "Прорыв трубы в подъезде",
        "easy",
        "14020300",
        31,
        "В подъезде жилого дома прорвало трубу горячей воды. Вода заливает первый этаж. Дежурная аварийная бригада доступна.",
        [
            "Бригада сообщает о выезде.",
            "Бригада прибыла. Выполняется перекрытие воды.",
            "Течь устранена, результаты работ переданы диспетчеру.",
        ],
        ["accepted", "responding", "completed"],
    ),
    (
        "elevator",
        "Застревание в лифте",
        "medium",
        "14100100",
        30,
        "Два взрослых человека застряли в лифте. Медицинская помощь не требуется. Дом обслуживает другая организация: учебная диспетчерская «Практика». Информация передана по принадлежности.",
        ["Укажите причину решения и факт передачи информации обслуживающей организации."],
        ["rejected"],
    ),
    (
        "wire",
        "Обрыв провода во дворе",
        "medium",
        "14110301",
        30,
        "Во дворе оборван провод неизвестного назначения. После приёма сообщения установлено: это линия связи другой организации. Информация передана в её диспетчерскую.",
        ["Отразите результаты уточнения и передачи информации в комментарии к статусу."],
        ["accepted", "refused"],
    ),
]
demo = []
for index, (slug, title, difficulty, code, page, prompt, briefing, expected) in enumerate(cases, 1):
    kind = by_code[code]
    card = {
        "caller_number": "+7000000000" + str(index),
        "name": "Учебный заявитель " + str(index),
        "address": f"Учебный город, Тренировочная улица, дом {index}",
        "incident_type_id": kind["id"],
        "comments": prompt,
    }
    demo.append(
        {
            "schema_version": "1.0",
            "id": slug,
            "title": title,
            "difficulty": difficulty,
            "service": "ДДС учебного района",
            "prompt": prompt,
            "briefing": briefing,
            "card": card,
            "source": {
                "type": "synthetic_adaptation",
                "file": "Работа с АРМ-112 для ДДС от ОКр_ГСИ.pdf",
                "page": page,
                "status": "TEAM_PROPOSAL",
                "note": "Синтетическая адаптация примера. Эталон требует подтверждения преподавателем.",
            },
            "reference": {
                "version": "1.0-demo",
                "address": card["address"],
                "incident_type_id": kind["id"],
                "expected_actions": expected,
            },
        }
    )
out = ROOT / "data_derived/scenarios"
out.mkdir(parents=True, exist_ok=True)
(out / "demo.json").write_text(json.dumps(demo, ensure_ascii=False, indent=2), encoding="utf-8")
